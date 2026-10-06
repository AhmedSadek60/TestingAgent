"""Document parsers (plug-ins registered in ``DOCUMENT_PARSERS``).

Each parser turns bytes into :class:`Parsed`: text blocks tagged with page/section, headings,
tables, metadata and *safety observations* (hidden text, active content, instruction-like text).
Parsers never execute document content. Where extraction is heuristic (PDF tables) the result is
flagged so reports do not overstate precision.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import Any

import yaml

from agentlab.core.errors import ParserError
from agentlab.core.plugins import Registry
from agentlab.documents.models import Heading, TableData
from agentlab.security import safeyaml

MAX_CHARS = 2_000_000


@dataclass
class Block:
    text: str
    page: int | None = None
    section: str | None = None
    kind: str = "paragraph"
    location: str = ""


@dataclass
class Parsed:
    blocks: list[Block] = field(default_factory=list)
    headings: list[Heading] = field(default_factory=list)
    tables: list[TableData] = field(default_factory=list)
    metadata: dict[str, Any] = field(default_factory=dict)
    pages: int | None = None
    warnings: list[str] = field(default_factory=list)
    hidden: list[str] = field(default_factory=list)
    active: list[str] = field(default_factory=list)
    unsupported: str | None = None


ParserFn = Callable[[bytes, str], Parsed]
DOCUMENT_PARSERS: Registry[ParserFn] = Registry("document_parsers")


def register(*extensions: str) -> Callable[[ParserFn], ParserFn]:
    def deco(fn: ParserFn) -> ParserFn:
        for ext in extensions:
            DOCUMENT_PARSERS.register(ext, fn, replace=True)
        return fn

    return deco


def _decode(data: bytes) -> str:
    if data[:3] == b"\xef\xbb\xbf":
        data = data[3:]
    return data[:MAX_CHARS].decode("utf-8", "replace")


def _paras(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]


# --------------------------------------------------------------------------------- plain text / markdown
@register(".txt", ".text", ".log", ".rst")
def parse_text(data: bytes, name: str) -> Parsed:
    text = _decode(data)
    out = Parsed()
    section: str | None = None
    for i, para in enumerate(_paras(text), 1):
        first = para.split("\n", 1)[0]
        if re.match(r"^([A-Z][A-Z0-9 \-:]{3,60}|\d+(\.\d+)*\s+[A-Z][\w \-]{2,60})$", first) and len(para) < 120:
            section = first.strip()
            out.headings.append(Heading(level=1, text=section))
        out.blocks.append(Block(para, section=section, location=f"paragraph {i}"))
    return out


@register(".md", ".markdown")
def parse_markdown(data: bytes, name: str) -> Parsed:
    text = _decode(data)
    out = Parsed()
    for m in re.finditer(r"<!--(.*?)-->", text, re.S):
        out.hidden.append(f"HTML comment: {m.group(1).strip()[:200]}")
    body = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    if body.startswith("---"):
        end = body.find("\n---", 3)
        if end > 0:
            try:
                out.metadata = safeyaml.load(body[3:end]) or {}
            except safeyaml.YamlRefused as exc:
                out.warnings.append(f"front matter was not read: {exc}")
            except yaml.YAMLError:
                pass
            body = body[end + 4 :]
    stack: list[str] = []
    in_code = False
    buf: list[str] = []
    start_line = 1
    lines = body.split("\n")

    def flush(kind: str = "paragraph") -> None:
        nonlocal buf
        if buf:
            txt = "\n".join(buf).strip()
            if txt:
                out.blocks.append(
                    Block(
                        txt,
                        section=" > ".join(stack) or None,
                        kind=kind,
                        location=f"lines {start_line}-{start_line + len(buf) - 1}",
                    )
                )
            buf = []

    for ln, line in enumerate(lines, 1):
        if line.strip().startswith("```"):
            if in_code:
                buf.append(line)
                flush("code")
                in_code = False
            else:
                flush()
                in_code = True
                start_line = ln
                buf.append(line)
            continue
        if in_code:
            buf.append(line)
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", line)
        if h:
            flush()
            level = len(h.group(1))
            stack[:] = stack[: level - 1] + [h.group(2).strip()]
            out.headings.append(Heading(level=level, text=h.group(2).strip()))
            out.blocks.append(
                Block(h.group(2).strip(), section=" > ".join(stack), kind="heading", location=f"line {ln}")
            )
            start_line = ln + 1
            continue
        if not line.strip():
            flush("table" if buf and all("|" in b for b in buf) else "paragraph")
            start_line = ln + 1
            continue
        if not buf:
            start_line = ln
        buf.append(line)
    flush()
    for b in out.blocks:
        if b.kind == "paragraph" and "|" in b.text and re.search(r"^\s*\|?\s*:?-{3,}", b.text, re.M):
            rows = [
                [c.strip() for c in r.strip().strip("|").split("|")]
                for r in b.text.split("\n")
                if r.strip() and not re.match(r"^\s*\|?\s*:?-{3,}", r)
            ]
            if rows:
                out.tables.append(TableData(section=b.section, header=rows[0], rows=rows[1:]))
                b.kind = "table"
    return out


# --------------------------------------------------------------------------------- csv / json / yaml
@register(".csv", ".tsv")
def parse_csv(data: bytes, name: str) -> Parsed:
    text = _decode(data)
    delim = "\t" if name.lower().endswith(".tsv") else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    out = Parsed()
    if not rows:
        return out
    header, body = rows[0], rows[1:]
    out.tables.append(TableData(header=header, rows=body[:5000]))
    for i, r in enumerate(body[:2000], 2):
        out.blocks.append(
            Block("; ".join(f"{h}: {v}" for h, v in zip(header, r, strict=False)), kind="row", location=f"row {i}")
        )
    return out


def _flatten(
    root: Any, *, max_items: int = 3000, max_visits: int = 100_000, max_depth: int = 40
) -> list[tuple[str, str]]:
    """``(path, value)`` pairs for the leaves of a parsed JSON or YAML value, in document order.

    The walk is bounded in the number of values it visits and the depth it descends, so a document that is small on disk
    and enormous in memory costs a fixed amount of work; whatever lies past the bounds is left out."""
    out: list[tuple[str, str]] = []
    stack: list[tuple[Any, str, int]] = [(root, "", 0)]
    visits = 0
    while stack and len(out) < max_items and visits < max_visits:
        node, path, depth = stack.pop()
        visits += 1
        if isinstance(node, dict | list):
            if depth >= max_depth:
                out.append((path, "(nested too deeply to read)"))
                continue
            if isinstance(node, dict):
                children = [(v, f"{path}.{k}" if path else str(k), depth + 1) for k, v in node.items()]
            else:
                children = [(v, f"{path}[{i}]", depth + 1) for i, v in enumerate(node[:500])]
            stack.extend(reversed(children))
        else:
            out.append((path, str(node)))
    return out


@register(".json")
def parse_json(data: bytes, name: str) -> Parsed:
    try:
        obj = json.loads(_decode(data))
    except (ValueError, RecursionError) as exc:
        raise ParserError(f"{name}: invalid JSON ({type(exc).__name__}: {str(exc)[:120]})") from exc
    out = Parsed(metadata={"top_level": type(obj).__name__})
    for path, val in _flatten(obj)[:3000]:
        if len(val) > 3:
            out.blocks.append(Block(f"{path}: {val}", kind="row", location=path, section=path.split(".")[0]))
    return out


@register(".yaml", ".yml")
def parse_yaml(data: bytes, name: str) -> Parsed:
    try:
        obj = safeyaml.load(_decode(data))
    except safeyaml.YamlRefused as exc:
        raise ParserError(f"{name}: not read ({exc})") from exc
    except yaml.YAMLError as exc:
        raise ParserError(f"{name}: invalid YAML ({exc})") from exc
    out = Parsed(metadata={"top_level": type(obj).__name__})
    for path, val in _flatten(obj)[:3000]:
        if len(val) > 3:
            out.blocks.append(Block(f"{path}: {val}", kind="row", location=path, section=path.split(".")[0]))
    return out


# --------------------------------------------------------------------------------- html
class _Html(HTMLParser):
    HIDDEN_STYLE = re.compile(
        r"display\s*:\s*none|visibility\s*:\s*hidden|font-size\s*:\s*0|opacity\s*:\s*0|"
        r"color\s*:\s*(#fff(?:fff)?|white)\b|height\s*:\s*0|left\s*:\s*-\d{3,}",
        re.I,
    )

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.blocks: list[Block] = []
        self.headings: list[Heading] = []
        self.hidden: list[str] = []
        self.tables: list[TableData] = []
        self.active: list[str] = []
        self._stack: list[tuple[str, bool]] = []
        self._buf: list[str] = []
        self._tag = ""
        self._section: list[str] = []
        self._row: list[str] = []
        self._rows: list[list[str]] = []
        self._in_table = False
        self._skip = 0

    def _hidden_attr(self, attrs: list[tuple[str, str | None]]) -> bool:
        d = {k: (v or "") for k, v in attrs}
        return "hidden" in d or d.get("aria-hidden") == "true" or bool(self.HIDDEN_STYLE.search(d.get("style", "")))

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        hid = self._hidden_attr(attrs)
        self._stack.append((tag, hid))
        if tag in {"script", "style", "noscript"}:
            self._skip += 1
            if tag == "script":
                self.active.append("<script> element present (not executed)")
        if tag == "table":
            self._in_table, self._rows = True, []
        if tag in {"iframe", "object", "embed", "form"}:
            self.active.append(f"<{tag}> element present")
        if tag in {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "td", "th", "div", "title"}:
            self._flush()
            self._tag = tag

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"}:
            self._skip = max(0, self._skip - 1)
        if tag in {"td", "th"}:
            self._row.append(" ".join("".join(self._buf).split()))
            self._buf = []
        elif tag == "tr":
            if self._row:
                self._rows.append(self._row)
            self._row = []
        elif tag == "table":
            if self._rows:
                self.tables.append(
                    TableData(section=" > ".join(self._section) or None, header=self._rows[0], rows=self._rows[1:])
                )
            self._in_table = False
        else:
            self._flush()
        while self._stack:
            t, _h = self._stack.pop()
            if t == tag:
                break

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        if any(h for _t, h in self._stack):
            if data.strip():
                self.hidden.append(f"hidden element text: {' '.join(data.split())[:200]}")
            return
        self._buf.append(data)

    def handle_comment(self, data: str) -> None:
        if data.strip():
            self.hidden.append(f"HTML comment: {data.strip()[:200]}")

    def _flush(self) -> None:
        text = " ".join("".join(self._buf).split())
        self._buf = []
        if not text or self._in_table:
            return
        if self._tag == "title":
            self.title = text
        elif self._tag.startswith("h") and len(self._tag) == 2 and self._tag[1].isdigit():
            lvl = int(self._tag[1])
            self._section[:] = self._section[: lvl - 1] + [text]
            self.headings.append(Heading(level=lvl, text=text))
            self.blocks.append(Block(text, section=" > ".join(self._section), kind="heading"))
        else:
            self.blocks.append(Block(text, section=" > ".join(self._section) or None))


@register(".html", ".htm")
def parse_html(data: bytes, name: str) -> Parsed:
    p = _Html()
    p.feed(_decode(data))
    p._flush()
    out = Parsed(
        blocks=p.blocks,
        headings=p.headings,
        tables=p.tables,
        hidden=p.hidden,
        active=p.active,
        metadata={"title": p.title},
    )
    for i, b in enumerate(out.blocks, 1):
        b.location = f"block {i}"
    return out


# --------------------------------------------------------------------------------- pdf
@register(".pdf")
def parse_pdf(data: bytes, name: str) -> Parsed:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except (PdfReadError, ValueError, OSError) as exc:
        raise ParserError(f"{name}: cannot read PDF ({type(exc).__name__}: {exc})") from exc
    out = Parsed(pages=len(reader.pages))
    if reader.is_encrypted:
        out.unsupported = "encrypted PDF: not decrypted (no password handling)"
        return out
    meta: dict[str, Any] = dict(reader.metadata or {})
    out.metadata = {k.lstrip("/").lower(): str(v)[:200] for k, v in meta.items() if v}
    root = reader.trailer.get("/Root", {})
    try:
        raw = bytes(data)
        for token in (b"/JavaScript", b"/JS ", b"/OpenAction", b"/Launch", b"/EmbeddedFile"):
            if token in raw:
                out.active.append(f"PDF contains {token.decode().strip('/ ')} (active content is never executed)")
        if root and "/AcroForm" in root:
            out.active.append("PDF contains a form (AcroForm)")
    except Exception:  # noqa: S110 - active-content scan is best-effort
        pass
    section: str | None = None
    for pno, page in enumerate(reader.pages, 1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:  # corrupted page
            out.warnings.append(f"page {pno}: text extraction failed ({type(exc).__name__})")
            continue
        if not text.strip():
            out.warnings.append(f"page {pno}: no extractable text (scanned image? OCR is not available)")
            continue
        groups: list[list[str]] = [[]]
        for line in text.split("\n"):
            stripped = line.strip()
            if not stripped:
                groups.append([])
            elif (
                re.match(r"^(\d+(\.\d+)*[.)]?\s+[A-Z][\w ,&\-]{2,70}|[A-Z][A-Z0-9 &\-]{4,60})$", stripped)
                and len(stripped) < 80
            ):
                groups.append([stripped])
                groups.append([])
            else:
                groups[-1].append(stripped)
        i = 0
        for g in groups:
            if not g:
                continue
            i += 1
            first = g[0]
            if (
                len(g) == 1
                and re.match(r"^(\d+(\.\d+)*[.)]?\s+[A-Z][\w ,&\-]{2,70}|[A-Z][A-Z0-9 &\-]{4,60})$", first)
                and len(first) < 80
            ):
                section = first
                out.headings.append(Heading(level=1, text=first, page=pno))
                out.blocks.append(
                    Block(first, page=pno, section=section, kind="heading", location=f"page {pno}, paragraph {i}")
                )
                continue
            # split the group into runs of tabular (multi-space / tab separated) and prose lines
            runs: list[tuple[bool, list[str]]] = []
            for ln in g:
                tab = bool(re.search(r"\S\s{2,}\S|\t", ln))
                if runs and runs[-1][0] == tab:
                    runs[-1][1].append(ln)
                else:
                    runs.append((tab, [ln]))
            for tab, run in runs:
                if tab and len(run) >= 3:
                    rows = [re.split(r"\s{2,}|\t", ln.strip()) for ln in run]
                    out.tables.append(
                        TableData(page=pno, section=section, header=rows[0], rows=rows[1:], heuristic=True)
                    )
                    out.blocks.append(
                        Block(
                            "\n".join(run),
                            page=pno,
                            section=section,
                            kind="table",
                            location=f"page {pno}, paragraph {i}",
                        )
                    )
                else:
                    out.blocks.append(
                        Block(" ".join(run), page=pno, section=section, location=f"page {pno}, paragraph {i}")
                    )
    if out.tables:
        out.warnings.append("PDF table extraction is heuristic (column alignment); verify table-derived facts manually")
    return out


# --------------------------------------------------------------------------------- docx
@register(".docx")
def parse_docx(data: bytes, name: str) -> Parsed:
    import docx  # python-docx

    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        if sum(i.file_size for i in zf.infolist()) > 200 * 1024 * 1024:
            raise ParserError(f"{name}: DOCX expands to an implausible size")
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise ParserError(f"{name}: cannot read DOCX ({type(exc).__name__}: {exc})") from exc
    out = Parsed()
    cp = document.core_properties
    out.metadata = {
        k: str(v)
        for k, v in {"title": cp.title, "author": cp.author, "created": cp.created, "modified": cp.modified}.items()
        if v
    }
    out.warnings.append("DOCX has no stored page numbers; page references are unavailable")
    stack: list[str] = []
    for i, para in enumerate(document.paragraphs, 1):
        text = para.text.strip()
        hidden_runs = [r.text for r in para.runs if r.font.hidden and r.text.strip()]
        for hr in hidden_runs:
            out.hidden.append(f"hidden text run: {hr[:200]}")
        if not text:
            continue
        style = (para.style.name or "") if para.style is not None else ""
        m = re.match(r"Heading (\d)", style)
        if m or style == "Title":
            lvl = int(m.group(1)) if m else 1
            stack[:] = stack[: lvl - 1] + [text]
            out.headings.append(Heading(level=lvl, text=text))
            out.blocks.append(Block(text, section=" > ".join(stack), kind="heading", location=f"paragraph {i}"))
        else:
            out.blocks.append(Block(text, section=" > ".join(stack) or None, location=f"paragraph {i}"))
    for t in document.tables:
        rows = [[c.text.strip() for c in r.cells] for r in t.rows]
        if rows:
            out.tables.append(TableData(section=" > ".join(stack) or None, header=rows[0], rows=rows[1:]))
            out.blocks.append(
                Block(
                    "\n".join(" | ".join(r) for r in rows),
                    kind="table",
                    section=" > ".join(stack) or None,
                    location="table",
                )
            )
    return out


# --------------------------------------------------------------------------------- source / images
_SOURCE_EXT = (
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".kt",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".cs",
    ".swift",
    ".c",
    ".cpp",
    ".h",
    ".sh",
    ".sql",
    ".toml",
    ".ini",
    ".cfg",
    ".xml",
)


@register(*_SOURCE_EXT)
def parse_source(data: bytes, name: str) -> Parsed:
    text = _decode(data)
    out = Parsed(metadata={"language_hint": name.rsplit(".", 1)[-1]})
    lines = text.split("\n")
    for start in range(0, len(lines), 40):
        chunk = "\n".join(lines[start : start + 40]).strip()
        if chunk:
            out.blocks.append(Block(chunk, kind="code", location=f"lines {start + 1}-{min(len(lines), start + 40)}"))
    return out


@register(".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp")
def parse_image(data: bytes, name: str) -> Parsed:
    out = Parsed(metadata={"bytes": len(data)})
    out.unsupported = (
        "image content (diagrams, screenshots, OCR) requires a multimodal provider; stored as an artifact only"
    )
    return out
