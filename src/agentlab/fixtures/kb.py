"""A small document knowledge base for the RAG fixtures: loaders, chunking and lexical retrieval.

It deliberately does not use AgentLab's own document analyser: the fixture is the *target*, and the platform that
evaluates it must not share its parsing code (a bug there would then hide in both).
"""

from __future__ import annotations

import csv
import functools
import html
import io
import logging
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

import docx
import pypdf
from docx.table import Table
from docx.text.paragraph import Paragraph

logging.getLogger("pypdf").setLevel(logging.ERROR)  # damaged files are expected here; pypdf's warnings are noise

DATA_DIR = Path(__file__).parent / "data"
RAG_DIR = DATA_DIR / "rag"
KB_FILES = ("hr-policy.pdf", "hr-policy-v1.md", "leave-policy.txt", "support-faq.docx", "catalog.csv")
RESTRICTED_FILES = ("executive-compensation.md",)  # in the index, but not for ordinary employees to read
SUPERSEDED = re.compile(r"\b(?:superseded|obsolete|deprecated|legacy|archived?|old)\b|[-_.]v\d\b", re.I)


def _words(text: str) -> frozenset[str]:
    return frozenset(text.split())


STOPWORDS = _words(
    "a an the of to in on at for and or is are was were be been by with from as it its this that these those what which who "
    "how many much do does did can could should would will if then than there their they we you your i my me about into "
    "per not no yes any all each every more most up down out over under"
)
# words that frame a question about the documents without being about their content
FRAME = _words(
    "according provided supplied documents document documentation say says said exact value figure source statement "
    "complete based only give tell state states sources comes come mention mentioned follow following first second two things need "
    "current rule regarding table column row page one sentence summarise summarize time long"
)
_TOKEN = re.compile(r"[a-z0-9]+(?:[-.'][a-z0-9]+)*")


# words a lexical index treats as one (a real retriever's embeddings would know)
SYNONYMS = {
    "delivery": "ship",
    "delivered": "ship",
    "shipping": "ship",
    "shipped": "ship",
    "ships": "ship",
    "keeping": "keep",
}


def stem(token: str) -> str:
    if token in SYNONYMS:
        return SYNONYMS[token]
    if token.endswith("ies") and len(token) > 4:
        return token[:-3] + "y"
    if token.endswith("s") and not token.endswith("ss") and len(token) > 3:
        return token[:-1]
    return token


def tokens(text: str, *, drop_frame: bool = False) -> list[str]:
    out = []
    for raw in _TOKEN.findall(text.lower().replace("$", "")):
        raw = raw.strip(".-'")
        if not raw or raw in STOPWORDS or (drop_frame and raw in FRAME):
            continue
        out.append(stem(raw))
    return out


@dataclass
class Chunk:
    source: str
    text: str
    page: int | None = None
    section: str = ""
    superseded: bool = False
    restricted: bool = False
    bag: set[str] = field(default_factory=set)

    def __post_init__(self) -> None:
        title = re.sub(r"[-_.]+", " ", Path(self.source).stem)  # retrievers index a passage with its document's title
        self.bag = set(tokens(f"{title} {self.section} {self.text}"))

    @property
    def lines(self) -> list[str]:
        return [ln.strip() for ln in self.text.splitlines() if ln.strip()]


def _sections(lines: list[str], heading: re.Pattern[str]) -> list[tuple[str, list[str]]]:
    out: list[tuple[str, list[str]]] = []
    current: tuple[str, list[str]] = ("", [])
    for ln in lines:
        if heading.match(ln):
            if current[1] or current[0]:
                out.append(current)
            current = (ln.strip(), [])
        else:
            current[1].append(ln)
    if current[1] or current[0]:
        out.append(current)
    return out


def pdf_chunks(name: str, data: bytes) -> list[Chunk]:
    chunks: list[Chunk] = []
    for number, page in enumerate(pypdf.PdfReader(io.BytesIO(data)).pages, 1):
        lines = [ln.rstrip() for ln in (page.extract_text() or "").splitlines() if ln.strip()]
        for title, body in _sections(lines, re.compile(r"^\d+\.\s+\S")):
            chunks.append(Chunk(name, "\n".join(body), page=number, section=title))
    return chunks


def markdown_chunks(name: str, body: str) -> list[Chunk]:
    old = bool(SUPERSEDED.search(body.splitlines()[0] if body.strip() else "")) or bool(
        SUPERSEDED.search(Path(name).stem)
    )
    return [
        Chunk(name, "\n".join(lines), section=title.lstrip("# ").strip(), superseded=old)
        for title, lines in _sections([ln for ln in body.splitlines() if ln.strip()], re.compile(r"^#{1,6}\s"))
        if lines
    ]


def text_chunks(name: str, body: str) -> list[Chunk]:
    old = bool(SUPERSEDED.search(Path(name).stem))
    return [
        Chunk(name, "\n".join(lines), section=title, superseded=old)
        for title, lines in _sections([ln for ln in body.splitlines() if ln.strip()], re.compile(r"^[A-Z][A-Z ]{3,}$"))
        if lines
    ]


def docx_chunks(name: str, data: bytes) -> list[Chunk]:
    document = docx.Document(io.BytesIO(data))
    chunks: list[Chunk] = []
    section = ""
    lines: list[str] = []

    def flush() -> None:
        if lines:
            chunks.append(Chunk(name, "\n".join(lines), section=section))

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            if not para.text.strip():
                continue
            if para.style is not None and para.style.name.lower().startswith(("heading", "title")):
                flush()
                section, lines = para.text.strip(), []
            else:
                lines.append(para.text.strip())
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    flush()
    return chunks


def csv_chunks(name: str, body: str) -> list[Chunk]:
    rows = list(csv.DictReader(io.StringIO(body)))
    return [Chunk(name, "; ".join(f"{k}: {v}" for k, v in row.items()), section=Path(name).stem) for row in rows]


def load_html(name: str, text: str) -> list[Chunk]:
    """A naive HTML-to-text pass: tags go, *all* text stays (hidden elements included), as in many real pipelines."""
    body = re.sub(r"(?is)<(?:script|style)\b.*?</(?:script|style)>", " ", text)
    body = re.sub(r"</(?:h\d|p|div|li|tr|title|section)>|<br\s*/?>", "\n", body, flags=re.I)
    body = html.unescape(re.sub(r"<[^>]+>", " ", body))
    lines = [" ".join(ln.split()) for ln in body.splitlines() if ln.strip()]
    title = re.search(r"(?is)<h1[^>]*>(.*?)</h1>", text)
    return [
        Chunk(
            name, "\n".join(lines), section=" ".join(re.sub(r"<[^>]+>", " ", title.group(1)).split()) if title else ""
        )
    ]


def chunks_from_text(name: str, text: str) -> list[Chunk]:
    """Chunks for text that arrives at run time (a document added to a session's knowledge)."""
    suffix = Path(name).suffix.lower()
    if suffix in {".html", ".htm"}:
        return load_html(name, text)
    if suffix == ".md":
        return markdown_chunks(name, text)
    if suffix == ".txt":
        return text_chunks(name, text)
    if suffix == ".csv":
        return csv_chunks(name, text)
    return [Chunk(name, text.strip())]


class UnreadableDocument(ValueError):
    """A file the reader cannot make sense of (damaged, empty, or not a kind it handles)."""


def chunks_from_bytes(name: str, data: bytes) -> list[Chunk]:
    """Chunks for a file's bytes. Raises :class:`UnreadableDocument` when it cannot be read."""
    suffix = Path(name).suffix.lower()
    try:
        if suffix == ".pdf":
            chunks = pdf_chunks(name, data)
        elif suffix == ".docx":
            chunks = docx_chunks(name, data)
        elif suffix in {".txt", ".md", ".csv", ".html", ".htm", ".py", ".json", ".yaml", ".yml"}:
            chunks = chunks_from_text(name, data.decode("utf-8"))
        else:
            raise UnreadableDocument(f"'{name}' is not a kind of document this reader handles")
    except UnreadableDocument:
        raise
    except Exception as exc:  # parsers raise a wide range of errors on damaged input
        raise UnreadableDocument(f"'{name}' could not be read ({type(exc).__name__})") from exc
    if not chunks:
        raise UnreadableDocument(f"'{name}' contains no readable text")
    return chunks


@functools.lru_cache(maxsize=1)
def default_chunks() -> tuple[Chunk, ...]:
    """The shipped knowledge base (parsed once per process). Treat the chunks as read-only."""
    chunks: list[Chunk] = []
    for name in KB_FILES:
        chunks.extend(chunks_from_bytes(name, (RAG_DIR / name).read_bytes()))
    for name in RESTRICTED_FILES:
        for chunk in chunks_from_bytes(name, (RAG_DIR / name).read_bytes()):
            chunk.restricted = True
            chunks.append(chunk)
    return tuple(chunks)


class KnowledgeBase:
    """Chunks plus IDF-weighted lexical retrieval."""

    def __init__(self, chunks: list[Chunk] | None = None) -> None:
        self.chunks: list[Chunk] = list(chunks or [])
        self.idf: dict[str, float] = {}
        self.unseen = 1.0
        self.reindex()

    @classmethod
    def default(cls, *, max_page: int | None = None) -> KnowledgeBase:
        """The shipped knowledge base; ``max_page`` drops everything after that page of a paged document."""
        return cls([c for c in default_chunks() if max_page is None or c.page is None or c.page <= max_page])

    def reindex(self) -> None:
        n = max(1, len(self.chunks))
        df: dict[str, int] = {}
        for c in self.chunks:
            for t in c.bag:
                df[t] = df.get(t, 0) + 1
        self.idf = {t: math.log(1 + n / d) for t, d in df.items()}
        self.unseen = math.log(1 + n)  # a word no chunk contains is at least as telling as the rarest one that is

    def with_chunks(self, extra: list[Chunk]) -> KnowledgeBase:
        return KnowledgeBase([*self.chunks, *extra])

    def cover(self, query: list[str], bag: set[str]) -> float:
        """Share of the query's information (IDF weight) that the words in ``bag`` contain, 0..1."""
        words = set(query)
        total = sum(self.idf.get(t, self.unseen) for t in words)
        if total <= 0:
            return 0.0
        return sum(self.idf.get(t, self.unseen) for t in words if t in bag) / total

    def coverage(self, query: list[str], chunk: Chunk) -> float:
        return self.cover(query, chunk.bag)

    def search(
        self,
        query: list[str],
        *,
        k: int = 3,
        old_bias: float = 0.6,
        include_restricted: bool = False,
        only_source: str | None = None,
    ) -> list[tuple[Chunk, float]]:
        """The ``k`` best chunks as ``(chunk, score)``. A superseded chunk's score is multiplied by ``old_bias``."""
        scored = []
        for c in self.chunks:
            if (c.restricted and not include_restricted) or (only_source and c.source.lower() != only_source.lower()):
                continue
            s = self.coverage(query, c)
            if s > 0:
                scored.append((c, s * (old_bias if c.superseded else 1.0)))
        scored.sort(key=lambda cs: -cs[1])
        return scored[:k]
